"""Score adapted retrieval JSON and write an Excel report.

Run from anywhere. Paths are resolved from the repository root.

``--folder`` is the chunk directory. Score files live under
``heavy_artifacts/evaluation/<folder>/``:

- ``<penalty>/<k>_<article_weight>_<sentence_weight>.json`` is an older cascade grid.
- ``<model>/<penalty>/<k>_<article_weight>_<sentence_weight>.json`` is a cascade grid for one embedder.
- ``<model>/dense/results.json``, ``<model>/rrf-k0/results.json``, and
  ``<model>/rrf-k60/results.json`` are the parallel baselines.

A cascade directory becomes one sheet, with a block per top-k. The three
parallel configs for one model share a sheet, one row each.
The workbook is written to ``rusbeir_eval/reports/<folder>.xlsx``.
"""

import argparse
import json
import sys
from pathlib import Path

import pandas as pd

SCRIPT_DIR = Path(__file__).resolve().parent
REPO_ROOT = SCRIPT_DIR.parent
RUSBEIR_DIR = REPO_ROOT / "rusbeir_eval"
if str(RUSBEIR_DIR) not in sys.path:
    sys.path.insert(0, str(RUSBEIR_DIR))

from rusBeIR.beir.datasets.data_loader_hf import HFDataLoader
from rusBeIR.beir.retrieval.evaluation import EvaluateRetrieval


FOLDER_TO_DATASET = {
    "sentences": "sents",
    "sw_2": "window_2",
    "sw_3": "window_3",
    "sw_4": "window_4",
    "sw_5": "window_5",
    "sw_6": "window_6",
    "para": "para",
}

PARALLEL_CONFIGS = ("dense", "rrf-k0", "rrf-k60")
MODEL_ORDER = ("e5", "frida", "arctic-l-v2", "giga-3b", "qwen3-4b")
METRIC_COLUMNS = ["NDCG@10", "MRR@10", "MAP@10", "Recall@10", "Recall@100"]
HEADLINE_METRICS = [
    ("NDCG@10", "ndcg", "NDCG@10"),
    ("MRR@10", "mrr", "MRR@10"),
    ("MAP@10", "map", "MAP@10"),
    ("Recall@10", "recall", "Recall@10"),
    ("Recall@100", "recall", "Recall@100"),
]


def dataset_repos(folder: str) -> tuple[str, str]:
    dataset_name = FOLDER_TO_DATASET[folder]
    return f"kaengreg/wikifacts-{dataset_name}", f"kaengreg/wikifacts-{dataset_name}-qrels"


def load_data(hf_repo, hf_repo_qrels, split="dev"):
    """Loads HF dataset and qrels."""
    print(f"Loading dataset {hf_repo} and qrels from {hf_repo_qrels}...")
    corpus, queries, qrels = HFDataLoader(
        hf_repo=hf_repo,
        hf_repo_qrels=hf_repo_qrels,
        streaming=False,
        keep_in_memory=False,
    ).load(split=split)
    return qrels


def score_results(results_path, qrels, retriever):
    """Return every BEIR metric dict for one adapted score file."""
    with open(results_path, "r", encoding="utf-8") as handle:
        results = json.load(handle)

    ndcg, average_precision, recall, precision = retriever.evaluate(qrels, results, retriever.k_values)
    mrr = retriever.evaluate_custom(qrels, results, retriever.k_values, "mrr")
    return {
        "ndcg": ndcg,
        "map": average_precision,
        "recall": recall,
        "precision": precision,
        "mrr": mrr,
    }


def headline_metrics(scored: dict) -> dict:
    return {label: scored[group].get(key, 0) for label, group, key in HEADLINE_METRICS}


def evaluate_single_file(results_path, qrels, retriever):
    """Evaluates a single JSON results file and returns the spreadsheet metrics."""
    return headline_metrics(score_results(results_path, qrels, retriever))


def parse_config_from_filename(filename):
    """Parses {k}_{w1}_{w2}.json into a readable config string."""
    stem = Path(filename).stem
    parts = stem.split("_")
    if len(parts) == 3 and parts[0].isdigit():
        k, w_art, w_sent = parts
        return f"weight_articles = {w_art}, weight_sentences = {w_sent}", k
    return stem, "unknown"


def is_grid_filename(filename: str) -> bool:
    stem = Path(filename).stem
    parts = stem.split("_")
    if len(parts) != 3 or not parts[0].isdigit():
        return False
    try:
        float(parts[1])
        float(parts[2])
    except ValueError:
        return False
    return True


def excel_sheet_name(name: str) -> str:
    cleaned = name.replace("/", "_").replace("\\", "_")
    for char in "[]:*?":
        cleaned = cleaned.replace(char, "_")
    return cleaned[:31] or "sheet"


def unique_sheet_name(name: str, used: set[str]) -> str:
    base = excel_sheet_name(name)
    candidate = base
    number = 2
    while candidate in used:
        suffix = f"_{number}"
        candidate = f"{base[:31 - len(suffix)]}{suffix}"
        number += 1
    used.add(candidate)
    return candidate


def discover_evaluations(base: Path):
    """Group score JSON files into cascade sheets and per-model baseline sheets.

    A directory of ``{k}_{w1}_{w2}.json`` files is one cascade sheet. Directories
    named ``dense``, ``rrf-k0``, or ``rrf-k60`` become rows on their model sheet.
    """
    run_dirs = []
    for path in [base, *base.rglob("*")]:
        if path.is_dir() and any(path.glob("*.json")):
            run_dirs.append(path)

    cascades = []
    baselines = {}
    for run_dir in sorted(run_dirs):
        files = sorted(path for path in run_dir.glob("*.json") if path.is_file())
        rel = run_dir.relative_to(base)
        rel_key = "results" if rel == Path(".") else rel.as_posix()

        if files and all(is_grid_filename(path.name) for path in files):
            cascades.append((rel_key, files))
            continue

        if rel != Path(".") and rel.name in PARALLEL_CONFIGS and rel.parent != Path("."):
            sheet_key = rel.parent.as_posix()
            row_label = rel.name
        elif rel != Path(".") and rel.name in PARALLEL_CONFIGS:
            sheet_key = "baselines"
            row_label = rel.name
        else:
            sheet_key = rel_key
            row_label = rel.name if rel != Path(".") else files[0].stem

        for json_file in files:
            label = row_label if len(files) == 1 else json_file.stem
            baselines.setdefault(sheet_key, []).append((label, json_file))

    return cascades, baselines


def _parallel_sort_key(label: str):
    if label in PARALLEL_CONFIGS:
        return (0, PARALLEL_CONFIGS.index(label), label)
    return (1, 0, label)


def _model_sort_key(sheet_key: str):
    model = sheet_key.split("/")[0]
    if model in MODEL_ORDER:
        return (0, MODEL_ORDER.index(model), sheet_key)
    return (1, 0, sheet_key)


def _style_table(worksheet, column_count: int, data_rows: int):
    from openpyxl.styles import Alignment, Font, PatternFill

    worksheet.column_dimensions["A"].width = 45
    for column_index in range(2, column_count + 1):
        worksheet.cell(row=1, column=column_index)
        letter = worksheet.cell(row=1, column=column_index).column_letter
        worksheet.column_dimensions[letter].width = 15

    header_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    for column_index in range(1, column_count + 1):
        header_cell = worksheet.cell(row=1, column=column_index)
        header_cell.fill = header_fill
        header_cell.font = Font(bold=True)
        header_cell.alignment = Alignment(horizontal="center", vertical="center")
        for row_offset in range(1, data_rows + 1):
            data_cell = worksheet.cell(row=1 + row_offset, column=column_index)
            data_cell.alignment = Alignment(horizontal="center", vertical="center")


def write_cascade_sheet(writer, sheet_name: str, json_files, qrels, retriever, output_excel: Path):
    from openpyxl.styles import Alignment, Font, PatternFill

    results_data = []
    for json_file in json_files:
        print(f"  Evaluating {json_file.name}...")
        metrics = evaluate_single_file(json_file, qrels, retriever)
        config_str, k_val = parse_config_from_filename(json_file.name)
        row = {"Config / Metrics": config_str, "k": k_val}
        row.update(metrics)
        results_data.append(row)

    df = pd.DataFrame(results_data)
    df["k_int"] = pd.to_numeric(df["k"], errors="coerce")
    df = df.sort_values(by=["k_int", "Config / Metrics"])
    cols = ["Config / Metrics", *METRIC_COLUMNS]

    if writer is None:
        for k_val in sorted(df["k_int"].unique()):
            print(f"\n{sheet_name} / Top-{int(k_val)} articles:")
            print(df[df["k_int"] == k_val][cols].to_string(index=False))
        return

    workbook = writer.book
    worksheet = workbook.create_sheet(sheet_name)
    writer.sheets[sheet_name] = worksheet
    worksheet.column_dimensions["A"].width = 45
    for col_letter in ["B", "C", "D", "E", "F"]:
        worksheet.column_dimensions[col_letter].width = 15

    current_row = 1
    for k_val in sorted(df["k_int"].unique()):
        df_k = df[df["k_int"] == k_val][cols]
        worksheet.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=len(cols))
        top_cell = worksheet.cell(row=current_row, column=1)
        top_cell.value = f"Top-{int(k_val)} articles"
        top_cell.alignment = Alignment(horizontal="center", vertical="center")
        top_cell.font = Font(bold=True)
        top_cell.fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
        current_row += 1

        df_k.to_excel(writer, sheet_name=sheet_name, startrow=current_row - 1, index=False)

        header_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
        for col_idx in range(1, len(cols) + 1):
            header_cell = worksheet.cell(row=current_row, column=col_idx)
            header_cell.fill = header_fill
            header_cell.alignment = Alignment(horizontal="center", vertical="center")
            for row_offset in range(1, len(df_k) + 1):
                data_cell = worksheet.cell(row=current_row + row_offset, column=col_idx)
                data_cell.alignment = Alignment(horizontal="center", vertical="center")
        current_row += len(df_k) + 2

    print(f"  Saved sheet '{sheet_name}' to {output_excel}")


def write_baseline_sheet(writer, sheet_name: str, rows, qrels, retriever, output_excel: Path):
    ordered = sorted(rows, key=lambda item: _parallel_sort_key(item[0]))
    results_data = []
    for label, json_file in ordered:
        print(f"  Evaluating {label}: {json_file.name}...")
        row = {"Config": label}
        row.update(evaluate_single_file(json_file, qrels, retriever))
        results_data.append(row)

    cols = ["Config", *METRIC_COLUMNS]
    df = pd.DataFrame(results_data)[cols]
    if writer is None:
        print(f"\n{sheet_name}:")
        print(df.to_string(index=False))
        return

    df.to_excel(writer, sheet_name=sheet_name, index=False)
    _style_table(writer.sheets[sheet_name], len(cols), len(df))
    print(f"  Saved sheet '{sheet_name}' to {output_excel}")


def main():
    parser = argparse.ArgumentParser(
        description=(
            "Score adapted retrieval JSON under heavy_artifacts/evaluation/<folder>/. "
            "Cascade grids keep one sheet per directory. "
            "dense, rrf-k0, and rrf-k60 share one sheet per embedder."
        )
    )
    parser.add_argument(
        "--folder",
        required=True,
        choices=sorted(FOLDER_TO_DATASET),
        help="Chunk directory inside heavy_artifacts/evaluation.",
    )
    args = parser.parse_args()

    hf_repo, hf_repo_qrels = dataset_repos(args.folder)
    base_input_dir = REPO_ROOT / "heavy_artifacts" / "evaluation" / args.folder
    output_excel = REPO_ROOT / "rusbeir_eval" / "reports" / f"{args.folder}.xlsx"
    output_excel.parent.mkdir(parents=True, exist_ok=True)

    if not base_input_dir.exists():
        print(f"Error: Input directory {base_input_dir} does not exist.")
        return

    cascades, baselines = discover_evaluations(base_input_dir)
    if not cascades and not baselines:
        print(f"No score JSON files found in {base_input_dir}.")
        return

    retriever = EvaluateRetrieval(k_values=[1, 3, 5, 10, 100])
    qrels = load_data(hf_repo, hf_repo_qrels)

    try:
        writer = pd.ExcelWriter(output_excel, engine="openpyxl")
    except ImportError:
        print("Warning: 'openpyxl' not found. Will only print results to console.")
        writer = None

    used_sheets = set()
    for rel_key, json_files in cascades:
        sheet_name = unique_sheet_name(rel_key, used_sheets)
        print(f"\nProcessing cascade: {rel_key}")
        write_cascade_sheet(writer, sheet_name, json_files, qrels, retriever, output_excel)

    for sheet_key in sorted(baselines, key=_model_sort_key):
        sheet_name = unique_sheet_name(sheet_key, used_sheets)
        print(f"\nProcessing baselines: {sheet_key}")
        write_baseline_sheet(writer, sheet_name, baselines[sheet_key], qrels, retriever, output_excel)

    if writer:
        workbook = writer.book
        if "Sheet" in workbook.sheetnames and len(workbook.sheetnames) > 1:
            workbook.remove(workbook["Sheet"])
        writer.close()
        print(f"\nBatch evaluation complete. Report saved to: {output_excel}")


if __name__ == "__main__":
    main()
